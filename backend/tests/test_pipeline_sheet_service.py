from datetime import date, datetime
from types import SimpleNamespace as NS
import uuid
from app.services import pipeline_sheet_service as s


def deal(**kw):
    base = dict(id=uuid.uuid4(), company="Acme", stage="warm", recurring_revenue=None, one_time_revenue=None,
                deal_value=None, win_probability=None, closure_eta=None, target_note=None,
                outcome_recorded_at=None, updated_at=datetime(2026, 10, 1), outcome_competitor=None)
    base.update(kw)
    return NS(**base)


def test_target_label_states():
    assert s.target_label("warm", date(2026, 10, 15), None) == "Oct 2026"
    assert s.target_label("warm", date(2026, 10, 15), "Oct/Nov") == "Oct/Nov"
    assert s.target_label("warm", None, None) == "—"
    assert s.target_label("on_hold", date(2026, 10, 1), None) == "Hold"
    assert s.target_label("closed_won", None, None, datetime(2026, 7, 9)) == "Closed in Jul"
    assert s.target_label("closed_lost", None, None, None, "Data Resolve") == "Lost to Data Resolve"
    assert s.target_label("closed_lost", None, None) == "Lost"


def test_probability_never_invented_for_open_deals():
    assert s.effective_probability("hot", None) is None
    assert s.effective_probability("hot", 70) == 70
    assert s.effective_probability("closed_won", None) == 100
    assert s.effective_probability("closed_lost", 90) == 0


def test_amounts_fallback_to_deal_value_for_old_deals():
    assert s.amounts(None, None, 500000) == {"arr": 0.0, "one_time": 0.0, "total": 500000.0}
    assert s.amounts(660000, 3500000, 1) == {"arr": 660000.0, "one_time": 3500000.0, "total": 4160000.0}
    assert s.amounts(None, 800000, None)["total"] == 800000.0


def test_row_weighted_and_history():
    r = s.build_row(deal(stage="hot", recurring_revenue=1000000, one_time_revenue=500000, win_probability=80,
                         closure_eta=date(2026, 11, 3)), ["Apr 2026", "Hold"])
    assert r["total"] == 1500000 and r["weighted"] == 1200000 and r["history"] == ["Apr 2026", "Hold"]
    assert r["target"] == "Nov 2026" and r["status"] == "open"


def test_group_by_rep_totals_exclude_lost_but_keep_won_and_hold():
    rows = [
        ("b", "Bela", s.build_row(deal(stage="hot", recurring_revenue=100, win_probability=50))),
        ("a", "Amit", s.build_row(deal(stage="closed_won", one_time_revenue=200))),
        ("a", "Amit", s.build_row(deal(stage="closed_lost", one_time_revenue=999))),
        ("a", "Amit", s.build_row(deal(stage="on_hold", recurring_revenue=50))),
    ]
    out = s.group_by_rep(rows)
    assert [g["rep_name"] for g in out["groups"]] == ["Amit", "Bela"]
    amit = out["groups"][0]["totals"]
    assert amit["count"] == 3 and amit["total"] == 250 and amit["weighted"] == 200
    assert out["totals"]["total"] == 350


def test_xlsx_export_builds_and_opens():
    from io import BytesIO
    from openpyxl import load_workbook
    from app.services.pipeline_sheet_export import build_xlsx
    rows = [("a", "Amit", s.build_row(deal(stage="hot", recurring_revenue=100, win_probability=70), ["Apr 2026"])),
            ("a", "Amit", s.build_row(deal(stage="closed_won", one_time_revenue=50)))]
    wb = load_workbook(BytesIO(build_xlsx(s.group_by_rep(rows))), rich_text=True)
    vals = [c.value for row in wb.active.iter_rows() for c in row if c.value is not None]
    assert "Grand total" in vals and any("Amit" in str(v) for v in vals)


def test_history_skips_current_and_repeats():
    r = s.build_row(deal(stage="hot", target_note="Jan-Feb 27"), ["Apr 2026", "Apr 2026", "Jan-Feb 27", "Hold"])
    assert r["history"] == ["Apr 2026", "Hold"]
