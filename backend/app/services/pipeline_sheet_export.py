"""Excel export of the rep-wise pipeline sheet. Reuses the MOM export's brand
colours and logos so every exported document looks like it came from one product."""
from io import BytesIO
from app.services.mom_export_service import BRAND_DARK, BRAND_LINE, BRAND_TINT, FLUIDPRO_LOGO, _logo

COLS = [("Pipeline", 30), ("ARR", 15), ("One time", 15), ("Total", 15), ("Win probability %", 14), ("Target month", 30)]
GREEN, GREY = "D9EAD3", "EEEEEE"


def _target_cell(row):
    """Struck-through earlier months followed by the current target, as one rich-text cell."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    parts = []
    for h in row["history"]:
        parts.append(TextBlock(InlineFont(strike=True, color="C00000"), h + "\n"))
    parts.append(row["target"])
    return CellRichText(*parts) if row["history"] else row["target"]


def build_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    dark = BRAND_DARK.lstrip("#")
    thin = Side(style="thin", color=BRAND_LINE.lstrip("#"))
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    inr = '"₹" #,##0'

    ws.row_dimensions[1].height = 42
    try:
        from openpyxl.drawing.image import Image as XlImage
        lg = _logo(FLUIDPRO_LOGO)
        if lg:
            img = XlImage(lg[0]); img.height, img.width = 34, int(34 * lg[1]); ws.add_image(img, "A1")
    except Exception:
        pass
    r = 3
    for i, (name, width) in enumerate(COLS, 1):
        c = ws.cell(r, i, name)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=dark)
        c.alignment, c.border = Alignment(horizontal="center", vertical="center", wrap_text=True), box
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "A4"

    def total_row(label, t, fill):
        nonlocal r
        r += 1
        vals = [label, t["arr"], t["one_time"], t["total"], None, f"Weighted ₹ {t['weighted']:,.0f}"]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.fill, c.border = Font(bold=True, color=dark), PatternFill("solid", fgColor=fill), box
            if i in (2, 3, 4):
                c.number_format = inr

    for g in data["groups"]:
        r += 1
        ws.cell(r, 1, f"{g['rep_name']}  ·  {g['totals']['count']} deals").font = Font(bold=True, size=12, color=dark)
        for i in range(1, 7):
            ws.cell(r, i).fill = PatternFill("solid", fgColor=BRAND_TINT.lstrip("#"))
        for row in g["rows"]:
            r += 1
            fill = GREEN if row["status"] == "won" else GREY if row["status"] == "lost" else None
            vals = [row["company"], row["arr"] or None, row["one_time"] or None, row["total"] or None,
                    (row["win_probability"] / 100) if row["win_probability"] is not None else None,
                    _target_cell(row)]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.border = box
                c.alignment = Alignment(vertical="top", wrap_text=(i in (1, 6)))
                if fill:
                    c.fill = PatternFill("solid", fgColor=fill)
                if i in (2, 3, 4):
                    c.number_format = inr
                if i == 5:
                    c.number_format = "0%"
                    c.alignment = Alignment(horizontal="center", vertical="top")
        total_row(f"Total — {g['rep_name']}", g["totals"], BRAND_LINE.lstrip("#"))
        r += 1
    total_row("Grand total", data["totals"], "FFD6E4")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
