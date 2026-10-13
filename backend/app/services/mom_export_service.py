"""Minutes of Meeting export — xlsx/pdf/docx.

One shared assembly step (build_minutes_document) + three thin format
renderers, so the three formats can't drift out of sync with each other
(see fluidgo-mom-architecture-lld.md ADR-3 for the library choices —
openpyxl/python-docx have no real alternative worth scoring; reportlab was
picked over weasyprint specifically to avoid adding system-level Cairo/
Pango dependencies to the Docker image).

"Best practice minutes" shape, same for all three formats: header
(company/date/purpose) -> attendees (Us/Customer, name+email) ->
discussion points (point/owner/date/status) -> AI summary text -> footer.
"""
from dataclasses import dataclass
from io import BytesIO
from typing import Optional
from app.config import settings


@dataclass
class MinutesDocument:
    company: str
    date: str
    purpose: str
    meeting_type: str
    attendees_us: list
    attendees_customer: list
    discussion_points: list
    ai_mom_summary: Optional[str]
    mom_status: str
    finalized_note: str


def build_minutes_document(m) -> MinutesDocument:
    """m is a Meeting ORM row or a dict from _serialize_meeting_detail —
    both are read the same way below."""
    def g(key, default=None):
        return m.get(key, default) if isinstance(m, dict) else getattr(m, key, default)

    attendees = g("attendees") or []
    return MinutesDocument(
        company=g("company") or "",
        date=str(g("date") or ""),
        purpose=(g("meeting_purpose") or "").replace("_", " ").title(),
        meeting_type=g("meeting_type") or "",
        attendees_us=[a for a in attendees if a.get("side") == "us"],
        attendees_customer=[a for a in attendees if a.get("side") == "customer"],
        discussion_points=g("discussion_points") or [],
        ai_mom_summary=g("ai_mom_summary"),
        mom_status=g("mom_status") or "none",
        finalized_note="Finalized" if g("mom_status") == "finalized" else "Draft, not yet finalized",
    )


# ── Presentation helpers (shared by every renderer) ──────────────────────────
import os
import re
from datetime import date as _date, datetime as _datetime
from html import escape as _esc

BRAND_DARK = "#1A0B2E"
BRAND_PINK = "#F0115E"
BRAND_MUTED = "#6B5F7A"
BRAND_TINT = "#F8F5FF"
BRAND_LINE = "#E8DFF5"


ASSETS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets")
FLUIDPRO_LOGO = os.path.join(ASSETS_DIR, "fluidpro_logo.png")  # portfolio brand  -> left of letterhead
WEPSOL_LOGO = os.path.join(ASSETS_DIR, "wepsol_logo.png")      # company brand    -> right of letterhead
# fluidGo is the *platform* name: it appears only in the footer, never as a header logo.


def _logo(path):
    """(path, width/height) or None when the file is missing/unreadable — a
    missing asset must degrade to a text-only header, never fail an export."""
    try:
        from PIL import Image
        with Image.open(path) as im:
            return path, im.size[0] / im.size[1]
    except Exception:
        return None


def fmt_date(value) -> str:
    """'2026-05-18' -> '18 May 2026'. Anything unparseable is returned as-is
    rather than raising — a bad date must never block an export."""
    if not value:
        return ""
    if isinstance(value, (_date, _datetime)):
        return f"{value.day} {value.strftime('%b %Y')}"
    s = str(value)[:10]
    try:
        d = _datetime.strptime(s, "%Y-%m-%d")
        return f"{d.day} {d.strftime('%b %Y')}"
    except ValueError:
        return str(value)


def tidy_name(name: str) -> str:
    """Capitalise names typed all-lowercase ('danish' -> 'Danish'). Names that
    already carry any capital (McDonald, A. K. Rao) are left exactly as typed."""
    name = (name or "").strip()
    return name.title() if name and name == name.lower() else name


def _split_attendee(a):
    return tidy_name(a.get("name", "")), (a.get("email") or "").strip()


def _attendee_lines(attendees):
    """Name, then the address in brackets only when one exists — no more
    '(no email)' noise in a customer-facing document."""
    out = []
    for a in attendees:
        name, email = _split_attendee(a)
        out.append(f"{name} ({email})" if email else name)
    return out or ["None listed"]


def _owner(p) -> str:
    side = (p.get("responsibility_side") or "").strip().title()
    name = tidy_name(p.get("responsibility_name") or "")
    if name and side:
        return f"{name} ({side})"
    return name or side or "Unassigned"


def _status(p) -> str:
    return (p.get("status") or "open").replace("_", " ").title()


def _target(p) -> str:
    return fmt_date(p.get("target_date")) or "TBD"


def _md_blocks(text: Optional[str]):
    """Tiny markdown subset -> [(kind, text)] where kind is 'h', 'li' or 'p'.
    The AI summary arrives as '## Heading' / '- bullet' / paragraphs; every
    format renders from this one parse so they look consistent."""
    blocks = []
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        m = re.match(r"^#{1,6}\s+(.*)$", line)
        if m:
            blocks.append(("h", m.group(1).strip()))
        elif re.match(r"^[-*•]\s+", line):
            blocks.append(("li", re.sub(r"^[-*•]\s+", "", line)))
        else:
            blocks.append(("p", line))
    return blocks


def _inline_html(text: str) -> str:
    """Escape, then turn **bold** into <b>. Safe for reportlab Paragraph and e-mail HTML."""
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", _esc(text))


def _plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def _file_stub(doc) -> str:
    return f"{doc.company} - {fmt_date(doc.date)}"


def _meta_rows(doc):
    return [
        ("Customer", doc.company or "-"),
        ("Date", fmt_date(doc.date) or "-"),
        ("Purpose", doc.purpose or "-"),
        ("Meeting type", doc.meeting_type or "-"),
    ]


# ── XLSX ─────────────────────────────────────────────────────────────────────
def render_xlsx(doc: MinutesDocument) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Minutes of Meeting"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    dark = BRAND_DARK.lstrip("#")
    head_fill = PatternFill("solid", fgColor=dark)
    head_font = Font(bold=True, color="FFFFFF")
    section_font = Font(bold=True, size=12, color=dark)
    thin = Side(style="thin", color=BRAND_LINE.lstrip("#"))
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.row_dimensions[1].height = 42  # white letterhead row: fluidPro left, wepsol right
    try:
        from openpyxl.drawing.image import Image as XlImage
        for name, anchor in ((FLUIDPRO_LOGO, "A1"), (WEPSOL_LOGO, "C1")):
            lg = _logo(name)
            if lg:
                img = XlImage(lg[0])
                img.height, img.width = 34, int(34 * lg[1])
                ws.add_image(img, anchor)
    except Exception:  # logos are decoration only
        pass
    ws.append([])
    ws.merge_cells("A2:D2")
    ws["A2"] = "Minutes of Meeting"
    ws["A2"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A2"].fill = head_fill
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 30
    ws.append([])
    ws.append([])

    for label, value in _meta_rows(doc):
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True, color=BRAND_MUTED.lstrip("#"))
    ws.append([])

    def section(title):
        ws.append([title])
        ws.cell(ws.max_row, 1).font = section_font

    def header(cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font, c.fill, c.border = head_font, head_fill, box

    section("Attendees")
    header(["Side", "Name", "Email", ""])
    rows = [("WEP Solutions", a) for a in doc.attendees_us] + [(doc.company or "Customer", a) for a in doc.attendees_customer]
    if rows:
        for side, a in rows:
            name, email = _split_attendee(a)
            ws.append([side, name, email or "-", ""])
            for c in ws[ws.max_row][:3]:
                c.border = box
    else:
        ws.append(["None listed"])
    ws.append([])

    section("Discussion Points")
    header(["Point", "Owner", "Target date", "Status"])
    if doc.discussion_points:
        for p in doc.discussion_points:
            ws.append([p.get("point", ""), _owner(p), _target(p), _status(p)])
            for c in ws[ws.max_row]:
                c.border = box
    else:
        ws.append(["None recorded"])
    ws.append([])

    if doc.ai_mom_summary:
        section("AI Minutes Summary")
        for kind, text in _md_blocks(doc.ai_mom_summary):
            if kind == "h":
                ws.append([_plain(text)])
                ws.cell(ws.max_row, 1).font = Font(bold=True)
            else:
                ws.merge_cells(start_row=ws.max_row + 1, start_column=1, end_row=ws.max_row + 1, end_column=4)
                ws.append([("• " if kind == "li" else "") + _plain(text)])
                ws.row_dimensions[ws.max_row].height = max(15, 15 * (len(text) // 95 + 1))
        ws.append([])

    ws.append([f"Generated by {settings.APP_NAME} | {doc.finalized_note}"])
    ws.cell(ws.max_row, 1).font = Font(size=9, color=BRAND_MUTED.lstrip("#"))

    for col, width in zip("ABCD", (60, 32, 26, 16)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── DOCX ─────────────────────────────────────────────────────────────────────
def _shade(cell, hex_fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill.lstrip("#"))
    tcPr.append(shd)


def render_docx(doc: MinutesDocument) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT

    d = Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE  # wide action table reads better landscape
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    base = d.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def run(par, text, bold=False, size=None, color=None):
        r = par.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        return r

    def heading(text):
        p = d.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(14), Pt(4)
        run(p, text, bold=True, size=13, color=BRAND_DARK)

    def table(rows, widths_cm, header=True):
        t = d.add_table(rows=0, cols=len(widths_cm))
        t.style = "Table Grid"
        t.autofit = False
        for ri, values in enumerate(rows):
            cells = t.add_row().cells
            for ci, (cell, val) in enumerate(zip(cells, values)):
                cell.width = Cm(widths_cm[ci])
                cell.text = ""
                par = cell.paragraphs[0]
                is_head = header and ri == 0
                run(par, str(val), bold=is_head, size=10, color="#FFFFFF" if is_head else None)
                if is_head:
                    _shade(cell, BRAND_DARK)
                elif ri % 2 == 0:
                    _shade(cell, BRAND_TINT)
        for ci, w in enumerate(widths_cm):  # LibreOffice/Word honour the column grid, not cell widths
            t.columns[ci].width = Cm(w)
        return t

    # Letterhead: fluidPro (portfolio) left, wepsol (company) right, on white
    lh = d.add_table(rows=1, cols=2)
    lh.autofit = False
    lcell, rcell = lh.rows[0].cells
    lh.columns[0].width, lh.columns[1].width = Cm(12.85), Cm(12.85)
    for c, name, align in ((lcell, FLUIDPRO_LOGO, WD_ALIGN_PARAGRAPH.LEFT), (rcell, WEPSOL_LOGO, WD_ALIGN_PARAGRAPH.RIGHT)):
        c.width = Cm(12.85)
        par = c.paragraphs[0]
        par.alignment = align
        lg = _logo(name)
        if lg:
            try:
                par.add_run().add_picture(lg[0], height=Cm(1.3))
            except Exception:
                pass
    # Title band
    band = d.add_table(rows=1, cols=1)
    cell = band.rows[0].cells[0]
    cell.width = Cm(25.7)
    _shade(cell, BRAND_DARK)
    p2 = cell.paragraphs[0]
    run(p2, "Minutes of Meeting", bold=True, size=20, color="#FFFFFF")
    p3 = cell.add_paragraph()
    run(p3, f"{doc.company}  |  {fmt_date(doc.date)}", size=10.5, color="#E8DFF5")

    heading("Meeting details")
    t = table(_meta_rows(doc), [4, 21.7], header=False)
    for row in t.rows:
        for r in row.cells[0].paragraphs[0].runs:
            r.bold = True

    heading("Attendees")
    rows = [("Side", "Name", "Email")]
    for a in doc.attendees_us:
        n, e = _split_attendee(a)
        rows.append(("WEP Solutions", n, e or "-"))
    for a in doc.attendees_customer:
        n, e = _split_attendee(a)
        rows.append((doc.company or "Customer", n, e or "-"))
    if len(rows) == 1:
        rows.append(("None listed", "", ""))
    table(rows, [7, 8, 10.7])

    heading("Discussion points")
    rows = [("#", "Point", "Owner", "Target date", "Status")]
    for i, p in enumerate(doc.discussion_points, 1):
        rows.append((i, p.get("point", ""), _owner(p), _target(p), _status(p)))
    if len(rows) == 1:
        rows.append(("-", "None recorded", "", "", ""))
    table(rows, [1, 11.5, 5.2, 3.8, 4.2])

    if doc.ai_mom_summary:
        heading("AI minutes summary")
        for kind, text in _md_blocks(doc.ai_mom_summary):
            if kind == "h":
                p = d.add_paragraph()
                p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(2)
                run(p, _plain(text), bold=True, size=11, color=BRAND_PINK)
            elif kind == "li":
                p = d.add_paragraph(style="List Bullet")
                run(p, _plain(text))
            else:
                run(d.add_paragraph(), _plain(text))

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(foot, f"Generated by {settings.APP_NAME}  |  {doc.finalized_note}", size=8.5, color=BRAND_MUTED)

    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────
def render_pdf(doc: MinutesDocument) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                                    ListFlowable, ListItem, KeepTogether, CondPageBreak)

    C = colors.HexColor
    PAGE = landscape(A4)
    page_w = PAGE[0] - 40 * mm  # 257 mm usable

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=14, textColor=C(BRAND_DARK), alignment=TA_LEFT)
    cell = ParagraphStyle("cell", parent=body, fontSize=9.5, leading=12.5)
    cell_head = ParagraphStyle("cellh", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    label = ParagraphStyle("label", parent=cell, fontName="Helvetica-Bold", textColor=C(BRAND_MUTED))
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=16, spaceAfter=6)
    h3 = ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=2, textColor=C(BRAND_PINK))

    def P(text, style=cell):
        return Paragraph(_inline_html(str(text)), style)

    def grid(data, widths, header=True):
        rows = []
        for ri, r in enumerate(data):
            rows.append([P(v, cell_head if (header and ri == 0) else cell) for v in r])
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        style = [
            ("GRID", (0, 0), (-1, -1), 0.5, C(BRAND_LINE)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), C(BRAND_DARK)),
                      ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, C(BRAND_TINT)])]
        t.setStyle(TableStyle(style))
        return t

    # Letterhead (white): fluidPro portfolio logo left, wepsol company logo right
    from reportlab.platypus import Image as RLImage
    logo_h = 12 * mm

    def logo_flow(name, align):
        lg = _logo(name)
        return RLImage(lg[0], width=logo_h * lg[1], height=logo_h, hAlign=align) if lg else ""

    letterhead = Table([[logo_flow(FLUIDPRO_LOGO, "LEFT"), logo_flow(WEPSOL_LOGO, "RIGHT")]],
                       colWidths=[page_w / 2, page_w / 2])
    letterhead.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (0, 0), "LEFT"), ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                                    ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 0)]))
    band = Table([[[Paragraph('<font color="#FFFFFF"><b>Minutes of Meeting</b></font>',
                              ParagraphStyle("b2", parent=body, fontSize=22, leading=26)),
                    Paragraph(f'<font color="#E8DFF5">{_esc(doc.company)}&nbsp;&nbsp;|&nbsp;&nbsp;{_esc(fmt_date(doc.date))}</font>',
                              ParagraphStyle("b3", parent=body, fontSize=11))]]], colWidths=[page_w])
    band.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C(BRAND_DARK)),
                              ("LEFTPADDING", (0, 0), (-1, -1), 14),
                              ("TOPPADDING", (0, 0), (-1, -1), 12), ("BOTTOMPADDING", (0, 0), (-1, -1), 12)]))

    story = [letterhead, band, Paragraph("Meeting details", h2)]
    meta = Table([[Paragraph(k, label), P(v)] for k, v in _meta_rows(doc)], colWidths=[35 * mm, page_w - 35 * mm])
    meta.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.5, C(BRAND_LINE)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                              ("LEFTPADDING", (0, 0), (-1, -1), 2)]))
    story.append(meta)

    story.append(Paragraph("Attendees", h2))
    att = [["Side", "Name", "Email"]]
    for a in doc.attendees_us:
        n, e = _split_attendee(a)
        att.append(["WEP Solutions", n, e or "-"])
    for a in doc.attendees_customer:
        n, e = _split_attendee(a)
        att.append([doc.company or "Customer", n, e or "-"])
    if len(att) == 1:
        att.append(["None listed", "", ""])
    story.append(grid(att, [60 * mm, 70 * mm, page_w - 130 * mm]))

    story.append(CondPageBreak(45 * mm))
    story.append(Paragraph("Discussion points", h2))
    dp = [["#", "Point", "Owner", "Target date", "Status"]]
    for i, p in enumerate(doc.discussion_points, 1):
        dp.append([i, p.get("point", ""), _owner(p), _target(p), _status(p)])
    if len(dp) == 1:
        dp.append(["-", "None recorded", "", "", ""])
    story.append(grid(dp, [10 * mm, 112 * mm, 55 * mm, 35 * mm, page_w - 212 * mm]))

    if doc.ai_mom_summary:
        story.append(CondPageBreak(45 * mm))
        story.append(Paragraph("AI minutes summary", h2))
        bullets = []

        def flush():
            if bullets:
                story.append(ListFlowable([ListItem(Paragraph(_inline_html(b), body), leftIndent=14) for b in bullets],
                                          bulletType="bullet", start="•", leftIndent=14, bulletFontSize=9))
                bullets.clear()

        for kind, text in _md_blocks(doc.ai_mom_summary):
            if kind == "li":
                bullets.append(text)
                continue
            flush()
            story.append(Paragraph(_inline_html(text), h3 if kind == "h" else body))
            if kind == "p":
                story.append(Spacer(1, 3))
        flush()

    footer_text = f"Generated by {settings.APP_NAME}  |  {doc.finalized_note}"

    def on_page(canvas, d):
        canvas.saveState()
        canvas.setStrokeColor(C(BRAND_LINE))
        canvas.line(20 * mm, 15 * mm, PAGE[0] - 20 * mm, 15 * mm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(C(BRAND_MUTED))
        canvas.drawString(20 * mm, 10 * mm, footer_text)
        canvas.drawRightString(PAGE[0] - 20 * mm, 10 * mm, f"Page {d.page}")
        canvas.restoreState()

    buf = BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=PAGE, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=12 * mm, bottomMargin=22 * mm,
                            title=f"Minutes of Meeting - {_file_stub(doc)}", author=settings.APP_NAME)
    pdf.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# ── E-mail body (HTML + plain text) ──────────────────────────────────────────
def render_email_html(doc: MinutesDocument) -> tuple:
    """Customer-facing e-mail body. Inline CSS only (mail clients strip <style>);
    solid colours instead of gradients (Outlook ignores gradients)."""
    def block(kind, text):
        t = _inline_html(text)
        if kind == "h":
            return f'<div style="font-weight:700;color:{BRAND_PINK};margin:14px 0 4px">{t}</div>'
        if kind == "li":
            return f'<div style="margin:2px 0 2px 14px">&bull;&nbsp;{t}</div>'
        return f'<p style="margin:6px 0">{t}</p>'

    summary_html = "".join(block(k, t) for k, t in _md_blocks(doc.ai_mom_summary)) or \
        '<p style="color:#6B5F7A">Please find the full minutes in the attached document.</p>'

    rows = "".join(
        f'<tr><td style="padding:6px 8px;border:1px solid {BRAND_LINE}">{_esc(p.get("point",""))}</td>'
        f'<td style="padding:6px 8px;border:1px solid {BRAND_LINE}">{_esc(_owner(p))}</td>'
        f'<td style="padding:6px 8px;border:1px solid {BRAND_LINE};white-space:nowrap">{_esc(_target(p))}</td>'
        f'<td style="padding:6px 8px;border:1px solid {BRAND_LINE}">{_esc(_status(p))}</td></tr>'
        for p in doc.discussion_points)
    table = (f'<div style="font-weight:700;margin:16px 0 6px">Action items</div>'
             f'<table style="border-collapse:collapse;width:100%;font-size:13px">'
             f'<tr style="background:{BRAND_DARK};color:#fff"><th align="left" style="padding:6px 8px">Point</th>'
             f'<th align="left" style="padding:6px 8px">Owner</th><th align="left" style="padding:6px 8px">Target date</th>'
             f'<th align="left" style="padding:6px 8px">Status</th></tr>{rows}</table>') if rows else ""

    html = f"""\
<div style="font-family:Segoe UI,Arial,sans-serif;max-width:620px;margin:0 auto;color:{BRAND_DARK}">
  <div style="background:{BRAND_DARK};padding:18px 22px;border-radius:10px 10px 0 0">
    <div style="color:{BRAND_PINK};font-size:13px;font-weight:700">fluidPro &nbsp;|&nbsp; WEP Solutions</div>
    <div style="color:#fff;font-size:20px;font-weight:700;margin-top:2px">Minutes of Meeting</div>
    <div style="color:#E8DFF5;font-size:13px;margin-top:2px">{_esc(doc.company)} &nbsp;|&nbsp; {_esc(fmt_date(doc.date))} &nbsp;|&nbsp; {_esc(doc.purpose)}</div>
  </div>
  <div style="border:1px solid {BRAND_LINE};border-top:none;padding:20px;border-radius:0 0 10px 10px;font-size:14px;line-height:1.5">
    <p style="margin:0 0 10px">Dear team,</p>
    <p style="margin:0 0 10px">Thank you for your time. Please find the minutes of our meeting below.</p>
    {summary_html}
    {table}
    <p style="color:#6B5F7A;font-size:11px;border-top:1px solid {BRAND_LINE};padding-top:10px;margin-top:18px">
      Sent via {_esc(settings.APP_NAME)}
    </p>
  </div>
</div>"""

    lines = [f"Minutes of Meeting - {doc.company} ({fmt_date(doc.date)})", ""]
    for kind, text in _md_blocks(doc.ai_mom_summary):
        lines.append(("- " if kind == "li" else "") + _plain(text))
    if doc.discussion_points:
        lines += ["", "Action items:"]
        lines += [f"- {p.get('point','')} | {_owner(p)} | {_target(p)} | {_status(p)}" for p in doc.discussion_points]
    lines += ["", f"Sent via {settings.APP_NAME}"]
    return html, "\n".join(lines)


# format -> (render_fn, mimetype, file_extension)
EXPORT_RENDERERS = {
    "xlsx": (render_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "docx": (render_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "pdf":  (render_pdf,  "application/pdf", "pdf"),
}
